"""Tests for resume text extraction and profile structuring."""

from __future__ import annotations

from io import BytesIO

import docx
import pytest
from fpdf import FPDF

from resume_agent import extract
from resume_agent.extract import ExtractionError, extract_profile, extract_text
from resume_agent.models import ExtractedProfile


def _make_pdf(text: str) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(0, 10, text)
    return bytes(pdf.output())


def _make_docx(*paragraphs: str) -> bytes:
    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    buf = BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_extract_text_pdf():
    text = extract_text(_make_pdf("Jane Doe Python Engineer"), "resume.pdf")
    assert "Jane Doe" in text


def test_extract_text_docx():
    text = extract_text(_make_docx("Jane Doe", "Python Engineer"), "resume.docx")
    assert "Jane Doe" in text
    assert "Python Engineer" in text


def test_extract_text_unsupported_type():
    with pytest.raises(ExtractionError, match="Unsupported file type"):
        extract_text(b"whatever", "resume.txt")


def test_extract_text_empty_document():
    with pytest.raises(ExtractionError, match="Could not read any text"):
        extract_text(_make_docx(), "resume.docx")  # no paragraphs


def test_extract_profile_delegates(monkeypatch):
    captured = {}
    expected = ExtractedProfile(name="Jane")

    def fake_complete(provider, model, system, user_message, schema, client):
        captured.update(
            provider=provider, model=model, schema=schema, user_message=user_message
        )
        return expected

    monkeypatch.setattr(extract, "complete_structured", fake_complete)

    result = extract_profile("raw resume text", "OpenAI", "gpt-4o")

    assert result is expected
    assert captured["provider"] == "OpenAI"
    assert captured["model"] == "gpt-4o"
    assert captured["schema"] is ExtractedProfile
    assert "raw resume text" in captured["user_message"]
